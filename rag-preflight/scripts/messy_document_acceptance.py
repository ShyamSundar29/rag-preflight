"""Generate and audit a bounded messy-document fixture without external files."""
from __future__ import annotations

import argparse
import base64
import io
import json
import logging
from pathlib import Path
import tempfile
from typing import Any

from rag_preflight import (check_source_folder, docx_receipt, pptx_receipt,
                           text_file_receipt)


_PNG = base64.b64decode(
    'iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII=')


def _image_only_pdf(path: Path) -> None:
    from pypdf import PdfWriter
    from pypdf.generic import (DecodedStreamObject, DictionaryObject, NameObject,
                               NumberObject)
    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    image = DecodedStreamObject()
    image.set_data(bytes([255, 255, 255]))
    image.update({NameObject('/Type'): NameObject('/XObject'),
                  NameObject('/Subtype'): NameObject('/Image'),
                  NameObject('/Width'): NumberObject(1),
                  NameObject('/Height'): NumberObject(1),
                  NameObject('/ColorSpace'): NameObject('/DeviceRGB'),
                  NameObject('/BitsPerComponent'): NumberObject(8)})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/XObject'):
        DictionaryObject({NameObject('/Im1'): writer._add_object(image)})})
    content = DecodedStreamObject()
    content.set_data(b'q 10 0 0 10 0 0 cm /Im1 Do Q')
    page[NameObject('/Contents')] = writer._add_object(content)
    with path.open('wb') as stream:
        writer.write(stream)


def _encrypted_pdf(path: Path) -> None:
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt('acceptance-only-password')
    with path.open('wb') as stream:
        writer.write(stream)


def _large_docx(path: Path) -> None:
    from docx import Document
    document = Document()
    document.add_heading('Messy document acceptance fixture', level=1)
    for number in range(1_000):
        document.add_paragraph(
            f'Policy paragraph {number}: retain independently inventoried source evidence.')
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Field'
    table.cell(0, 1).text = 'Value'
    table.cell(1, 0).text = 'coverage'
    table.cell(1, 1).text = 'required'
    document.save(path)


def _image_heavy_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches
    presentation = Presentation()
    for number in range(8):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(io.BytesIO(_PNG), Inches(1), Inches(1),
                                 width=Inches(2), height=Inches(2))
        if number == 0:
            box = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(6), Inches(1))
            box.text = 'Only this slide contains extractable text; images are not OCRed.'
    presentation.save(path)


def build_fixture(root: Path) -> Path:
    """Create deterministic edge cases; inputs are synthetic and license-safe."""
    root.mkdir(parents=True, exist_ok=True)
    (root / 'clean.txt').write_text('A valid UTF-8 policy record.', encoding='utf-8')
    (root / 'invalid-utf8.txt').write_bytes(b'not utf-8: \xff')
    (root / 'damaged.pdf').write_bytes(b'%PDF-1.7\nnot a valid PDF')
    (root / 'empty.pdf').write_bytes(b'')
    (root / 'useful.html').write_text(
        '<html><body><h1>Visible policy</h1><script>ignore me</script>'
        '<p>Retain this text.</p></body></html>', encoding='utf-8')
    (root / 'script-only.html').write_text(
        '<script>not searchable content</script>', encoding='utf-8')
    (root / 'unsupported.csv').write_text('not,part,of,the,document,subset\n', encoding='utf-8')
    _image_only_pdf(root / 'image-only.pdf')
    _encrypted_pdf(root / 'encrypted.pdf')
    _large_docx(root / 'large.docx')
    _image_heavy_pptx(root / 'image-heavy.pptx')
    expected = root.parent / 'expected-files.txt'
    names = sorted(path.name for path in root.iterdir()
                   if path.suffix.lower() in ('.pdf', '.txt', '.html', '.docx', '.pptx'))
    expected.write_text('\n'.join(names + ['missing-policy.pdf']) + '\n', encoding='utf-8')
    return expected


def run_acceptance() -> dict[str, Any]:
    with tempfile.TemporaryDirectory(prefix='rag-preflight-messy-') as folder:
        root = Path(folder) / 'sources'
        expected = build_fixture(root)
        # Keep the acceptance artifact machine-readable; the compact report
        # retains the categorized parser failure without pypdf's raw log line.
        pypdf_logger = logging.getLogger('pypdf')
        old_level = pypdf_logger.level
        pypdf_logger.setLevel(logging.CRITICAL + 1)
        try:
            report = check_source_folder(root, expected=expected, max_examples=10)
        finally:
            pypdf_logger.setLevel(old_level)
        groups = {group['code']: group for group in report.by_code(max_examples=10)}
        required = {'source_read_failed', 'likely_image_only_page',
                    'missing_expected_source'}
        if not required.issubset(groups):
            raise AssertionError(f'Missing expected findings: {sorted(required - groups.keys())}')
        if report.files_discovered != 10 or report.files_audited != 6:
            raise AssertionError('Unexpected supported-file discovery or processing count')
        if report.unsupported_files != 1 or report.image_only_units != 1:
            raise AssertionError('Unexpected unsupported-file or image-only count')
        reasons = {sample.get('reason') for sample in groups['source_read_failed']['examples']}
        if not {'invalid_utf8', 'empty_file', 'invalid_pdf', 'encrypted_pdf'}.issubset(reasons):
            raise AssertionError(f'Read-failure categories were not preserved: {sorted(reasons)}')
        word = docx_receipt(root / 'large.docx', document_id='large.docx')
        slides = pptx_receipt(root / 'image-heavy.pptx', document_id='image-heavy.pptx')
        html = text_file_receipt(root / 'useful.html', document_id='useful.html')
        observations = {
            'invalid_utf8_rejected': 'invalid_utf8' in reasons,
            'empty_and_damaged_pdf_rejected': {'empty_file', 'invalid_pdf'}.issubset(reasons),
            'encrypted_pdf_rejected_without_password': 'encrypted_pdf' in reasons,
            'image_only_pdf_flagged_for_possible_ocr': report.image_only_units == 1,
            'large_docx_processed': ('Policy paragraph 999' in word.texts[0][1]
                                      and 'coverage' in word.texts[0][1]),
            'image_heavy_pptx_empty_slides_visible': (
                slides.document.expected_units == tuple(f'slide:{n}' for n in range(1, 9))
                and slides.receipt.empty_units == tuple(f'slide:{n}' for n in range(2, 9))),
            'static_html_excludes_script': ('Retain this text' in html.texts[0][1]
                                            and 'ignore me' not in html.texts[0][1]),
            'missing_expected_file_detected': report.missing_expected_files == 1,
            'unsupported_file_counted': report.unsupported_files == 1,
        }
        if not all(observations.values()):
            raise AssertionError(f'Acceptance observation failed: {observations}')
        return {
            'schema_version': 1,
            'fixture_kind': 'generated_synthetic_edge_cases',
            'outcome': report.outcome,
            'metrics': report.metrics(),
            'finding_groups': report.by_code(max_examples=3),
            'read_failure_reasons': sorted(reasons),
            'expected_observations': observations,
            'limitations': [
                'Fixtures are generated edge cases, not evidence from an external production corpus.',
                'Image presence is a heuristic; the library does not perform OCR.',
                'DOCX and PPTX helpers intentionally cover only their documented text subsets.',
            ],
        }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, help='Optional path for the JSON result')
    args = parser.parse_args(argv)
    result = run_acceptance()
    rendered = json.dumps(result, indent=2, sort_keys=True) + '\n'
    if args.output is not None:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(rendered, encoding='utf-8')
    print(rendered, end='')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
