"""Extraction at source-unit boundaries; receipts never inferred from chunks."""
from collections.abc import Callable, Iterable
from dataclasses import dataclass
import hashlib
from io import BytesIO
from pathlib import Path
from typing import Any

from .ingestion import DocumentSpec, ExtractionReceipt


@dataclass(frozen=True)
class ExtractionResult:
    document: DocumentSpec
    receipt: ExtractionReceipt
    texts: tuple[tuple[str, str], ...]
    failures: tuple[tuple[str, str], ...]
    likely_image_only_units: tuple[str, ...] = ()
    image_inspection_failed_units: tuple[str, ...] = ()

    def chunks(self, splitter: Callable[[str], Iterable[str]] | None = None, *, source: str) -> list[dict[str, Any]]:
        """Split inside each unit; context may be cut at page boundaries. Audit afterward."""
        from .adapters import assign_chunk_keys
        result = []
        for unit, text in self.texts:
            if not text.strip():
                continue
            parts = (text,) if splitter is None else splitter(text)
            if isinstance(parts, str):
                raise TypeError('splitter must return an iterable of strings, not a string')
            for part in parts:
                if not isinstance(part, str):
                    raise TypeError('splitter outputs must be strings')
                result.append({'text': part, 'metadata': {'document_id': self.document.document_id,
                    'source_version': self.document.source_version, 'source': source, 'units': [unit]}})
        return assign_chunk_keys(result)


def receipt_from_callable(document: DocumentSpec, extract_unit: Callable[[str], str | None]) -> ExtractionResult:
    """Attempt every independently declared unit. None/blank is empty, exceptions failed.

    Inventory must come from the source, not the successful extractor output.
    Ordinary exceptions are recorded by type without potentially sensitive messages.
    Cancellation and process interrupts propagate. No retries or OCR are performed.
    """
    if document.expected_units is None:
        raise ValueError('An independent expected-unit inventory is required')
    texts, failed = [], []
    for unit in document.expected_units:
        assert isinstance(unit, str)  # DocumentSpec canonicalizes integer aliases.
        try:
            text = extract_unit(unit)
            if text is None:
                text = ''
            if not isinstance(text, str):
                raise TypeError('Extractor must return text or None')
            texts.append((unit, text))
        except Exception as exc:
            failed.append((unit, type(exc).__name__))
    receipt = ExtractionReceipt(document.document_id, document.source_version,
        processed_units=tuple(u for u, _ in texts), empty_units=tuple(u for u, t in texts if not t.strip()),
        failed_units=tuple(u for u, _ in failed), completed=True)
    return ExtractionResult(document, receipt, tuple(texts), tuple(failed))


def pypdf_receipt(path: str | Path, *, document_id: str,
                  allowed_empty_pages: Iterable[int] = ()) -> ExtractionResult:
    """Read one immutable byte copy, inventory its page tree, then attempt each page.

    source_version is SHA-256 of the bytes. Page-tree/open failures raise because
    an honest inventory cannot be established. Install rag-preflight[pdf].
    """
    from pypdf import PdfReader
    raw = Path(path).read_bytes()
    reader = PdfReader(BytesIO(raw))
    count = len(reader.pages)
    document = DocumentSpec(document_id, hashlib.sha256(raw).hexdigest(),
        expected_units=tuple(range(1, count + 1)), allowed_empty_units=tuple(allowed_empty_pages),
        min_chunks=0 if count == 0 else 1)
    result = receipt_from_callable(document, lambda unit: reader.pages[int(unit.split(':')[1]) - 1].extract_text())
    image_only: list[str] = []
    inspection_failed: list[str] = []
    for unit in result.receipt.empty_units:
        assert isinstance(unit, str)  # Receipt canonicalizes the page aliases.
        try:
            if len(reader.pages[int(unit.split(':')[1]) - 1].images) > 0:
                image_only.append(unit)
        except Exception:
            inspection_failed.append(unit)
    return ExtractionResult(result.document, result.receipt, result.texts, result.failures,
                            tuple(image_only), tuple(inspection_failed))


def text_file_receipt(path: str | Path, *, document_id: str) -> ExtractionResult:
    """UTF-8 .txt/.md/.rst and static .html/.htm body text; no rendered JavaScript."""
    from html.parser import HTMLParser
    raw = Path(path).read_bytes()
    text = raw.decode('utf-8')
    if Path(path).suffix.lower() in ('.html', '.htm'):
        class BodyText(HTMLParser):
            def __init__(self) -> None:
                super().__init__(convert_charrefs=True)
                self.hidden = 0
                self.parts: list[str] = []
            def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
                if tag in ('script', 'style', 'template'):
                    self.hidden += 1
            def handle_endtag(self, tag: str) -> None:
                if tag in ('script', 'style', 'template') and self.hidden:
                    self.hidden -= 1
            def handle_data(self, data: str) -> None:
                if not self.hidden:
                    self.parts.append(data)
        parser = BodyText()
        parser.feed(text)
        parser.close()
        text = ' '.join(parser.parts)
    document = DocumentSpec(document_id, hashlib.sha256(raw).hexdigest(), expected_units=('file',))
    return receipt_from_callable(document, lambda unit: text)


def docx_receipt(path: str | Path, *, document_id: str) -> ExtractionResult:
    """Optional python-docx: body paragraphs and table cells, one file-level unit.

    Headers, footers, floating text boxes, images, tracked changes and embedded
    objects are outside this subset. File-level coverage cannot detect dropped
    paragraphs in a nonempty body. Install rag-preflight[office].
    """
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    raw = Path(path).read_bytes()
    document = DocumentSpec(document_id, hashlib.sha256(raw).hexdigest(), expected_units=('body',))
    source = Document(BytesIO(raw))
    parts: list[str] = []
    for item in source.iter_inner_content():
        if isinstance(item, Paragraph):
            parts.append(item.text)
        elif isinstance(item, Table):
            for row in item.rows:
                for cell in row.cells:
                    parts.extend(p.text for p in cell.paragraphs)
    return receipt_from_callable(document, lambda unit: '\n'.join(parts))


def pptx_receipt(path: str | Path, *, document_id: str) -> ExtractionResult:
    """Optional python-pptx: independently enumerate slides, extract visible shape text.

    Speaker notes, images, charts and embedded objects are not searchable text.
    Slide representation does not prove all within-slide content was extracted.
    Install rag-preflight[office].
    """
    from pptx import Presentation
    raw = Path(path).read_bytes()
    slides = Presentation(BytesIO(raw)).slides
    document = DocumentSpec(document_id, hashlib.sha256(raw).hexdigest(),
                            expected_units=tuple(f'slide:{n}' for n in range(1, len(slides) + 1)))
    def shape_text(shapes: Any) -> list[str]:
        parts: list[str] = []
        for shape in shapes:
            if shape.has_text_frame:
                parts.append(shape.text)
            if shape.has_table:
                parts.extend(cell.text for row in shape.table.rows for cell in row.cells)
            if shape.shape_type == 6:  # GROUP; recurse into nested text shapes.
                parts.extend(shape_text(shape.shapes))
        return parts
    return receipt_from_callable(document,
        lambda unit: '\n'.join(shape_text(slides[int(unit.split(':')[1]) - 1].shapes)))
