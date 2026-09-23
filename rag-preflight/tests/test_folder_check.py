"""Real file-format and CLI evidence for the current-source folder door."""
import json
import io
from pathlib import Path
import tempfile
import unittest
from unittest import mock

from rag_preflight import (check_source_folder, docx_receipt, pptx_receipt,
                           text_file_receipt)
from rag_preflight.cli import main
from rag_preflight.folder import _read_failure


class FolderCheckTests(unittest.TestCase):
    def test_empty_scope_is_setup_error_without_none_prefix(self):
        with tempfile.TemporaryDirectory() as tmp:
            report = check_source_folder(tmp)
            self.assertEqual(report.outcome, 'no_supported_files')
            self.assertFalse(report.passed)
            self.assertNotIn('None:', report.table())
            with mock.patch('sys.stdout', new_callable=io.StringIO) as output:
                self.assertEqual(main(['check', tmp]), 2)
            self.assertIn('NO SUPPORTED FILES', output.getvalue())

    def test_missing_optional_reader_and_bounded_examples(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for n in range(12):
                (root / f'{n}.pdf').write_bytes(b'%PDF-1.4')
            with mock.patch('rag_preflight.folder.pypdf_receipt', side_effect=ImportError('pypdf')):
                report = check_source_folder(root, max_examples=2)
                self.assertEqual(report.outcome, 'reader_unavailable')
                self.assertEqual(report.reader_unavailable_files, 12)
                self.assertEqual(len(report.by_code(max_examples=2)[0]['examples']), 2)
                self.assertNotIn('empty_inventory', {g['code'] for g in report.by_code()})
                self.assertIn('rag-preflight[pdf]', report.table())
                with mock.patch('sys.stdout', new_callable=io.StringIO):
                    self.assertEqual(main(['check', str(root)]), 2)

    def test_expected_list_detects_absence_and_rejects_bad_paths(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / 'a').mkdir()
            (root / 'a/present.txt').write_text('Present.')
            listing = root / 'expected.list'
            listing.write_text('a/present.txt\na/missing.txt\n')
            report = check_source_folder(root, expected=listing)
            self.assertEqual(report.outcome, 'audit_failed')
            self.assertEqual(report.expected_files, 2)
            self.assertEqual(report.missing_expected_files, 1)
            self.assertIn('a/missing.txt', report.table())
            self.assertNotIn('Unverified: expected files absent from this folder', report.table())
            self.assertNotIn('expected_files_outside_current_folder', report.to_dict()['checks_unverified'])
            with mock.patch('sys.stdout', new_callable=io.StringIO):
                self.assertEqual(main(['check', str(root), '--expected', str(listing)]), 1)
            for invalid in ('../escape.txt', '/absolute.txt', 'a//present.txt',
                            'a/present.txt\na/present.txt', 'ignored.bin'):
                listing.write_text(invalid)
                with self.assertRaises(ValueError):
                    check_source_folder(root, expected=listing)

    def test_plain_language_read_errors_and_partial_failure(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / 'good.txt').write_text('Good text.')
            (root / 'bad.txt').write_bytes(b'\xff')
            report = check_source_folder(root)
            self.assertEqual(report.outcome, 'incomplete')
            self.assertIn('not valid UTF-8', report.table())
            self.assertEqual(report.files_audited, 1)
            with mock.patch('sys.stdout', new_callable=io.StringIO):
                self.assertEqual(main(['check', str(root)]), 2)

    def test_pdf_error_categories_and_expected_scope_under_exclusion(self):
        PdfStreamError = type('PdfStreamError', (Exception,), {})
        EmptyFileError = type('EmptyFileError', (Exception,), {})
        self.assertEqual(_read_failure(Path('bad.pdf'), PdfStreamError())['reason'], 'invalid_pdf')
        self.assertEqual(_read_failure(Path('empty.pdf'), EmptyFileError())['reason'], 'empty_file')
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / 'sources'
            root.mkdir()
            listing = Path(tmp) / 'expected.txt'
            listing.write_text('missing.txt\n')
            (root / 'linked').symlink_to(Path(tmp), target_is_directory=True)
            report = check_source_folder(root, expected=listing)
            self.assertEqual(report.outcome, 'incomplete')
            self.assertEqual(report.missing_expected_files, 0)
            self.assertIn('expected_file_absence', report.to_dict()['checks_unverified'])

    def test_nested_text_html_and_plain_language_missing_unit(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / 'a').mkdir()
            (root / 'b').mkdir()
            (root / 'a/report.txt').write_text('Useful source text. ' * 40)
            (root / 'b/report.txt').write_text('Different useful source. ' * 40)
            (root / 'page.html').write_text('<h1>Useful HTML</h1><script>secret JS</script><p>Body text.</p>')
            report = check_source_folder(root)
            self.assertTrue(report.passed, report.table())
            self.assertEqual(report.files_discovered, 3)
            self.assertEqual({f.document_id for f in report.audit.findings}, set())
            self.assertEqual(report.metrics()['files_audited'], 3)
            self.assertEqual(report.metrics()['missing_units'], 0)
            self.assertNotIn('secret JS', str(report.to_dict()))
            self.assertNotIn('secret JS', text_file_receipt(root / 'page.html',
                document_id='page.html').texts[0][1])
            self.assertIn('historical_index_coverage', report.to_dict()['checks_unverified'])
            (root / 'blank.html').write_text('<script>only code</script>')
            broken = check_source_folder(root)
            self.assertFalse(broken.passed)
            self.assertIn('blank.html', broken.table())
            self.assertIn('no extractable text', broken.table().lower())
            self.assertTrue(broken.audit.findings)  # detailed audit remains accessible

    def test_symlink_and_bad_utf8_are_not_silently_omitted(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / 'root'
            root.mkdir()
            (root / 'good.txt').write_text('Good text.')
            (root / 'bad.txt').write_bytes(b'\xff')
            (root / 'linked.txt').symlink_to(Path(tmp) / 'outside.txt')
            report = check_source_folder(root)
            self.assertFalse(report.passed)
            codes = {g['code'] for g in report.by_code()}
            self.assertIn('source_read_failed', codes)
            self.assertIn('source_path_excluded', codes)
            self.assertEqual(report.files_discovered, 3)
            self.assertEqual(report.files_audited, 1)
            self.assertNotIn('good.txt', str(report.metrics()['findings']))

    def test_optional_office_body_and_slide_units(self):
        try:
            from docx import Document
            from pptx import Presentation
        except ImportError:
            self.skipTest('rag-preflight[office] is unavailable')
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            word = Document()
            word.add_paragraph('A paragraph that should be represented.')
            word.add_table(rows=1, cols=1).cell(0, 0).text = 'A table cell that should be represented.'
            word.save(root / 'manual.docx')
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[5])
            box = slide.shapes.add_textbox(1, 1, 100, 100)
            box.text = 'The slide text should be represented.'
            deck.save(root / 'brief.pptx')
            report = check_source_folder(root)
            self.assertTrue(report.passed, report.table())
            self.assertEqual(report.units_enumerated, 2)
            self.assertEqual(report.files_audited, 2)
            self.assertIn('table cell', docx_receipt(root / 'manual.docx',
                document_id='manual.docx').texts[0][1])
            self.assertIn('slide text', pptx_receipt(root / 'brief.pptx',
                document_id='brief.pptx').texts[0][1])
            self.assertIn('within_unit_text_completeness', report.to_dict()['checks_unverified'])
            blank = Presentation()
            blank.slides.add_slide(blank.slide_layouts[6])
            blank.save(root / 'blank.pptx')
            broken = check_source_folder(root)
            self.assertFalse(broken.passed)
            self.assertIn('blank.pptx', broken.table())

    def test_real_pdf_image_only_page_is_a_heuristic_warning(self):
        try:
            from pypdf import PdfWriter
            from pypdf.generic import (DecodedStreamObject, DictionaryObject,
                                       NameObject, NumberObject)
        except ImportError:
            self.skipTest('rag-preflight[pdf] is unavailable')
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            writer = PdfWriter()
            page = writer.add_blank_page(width=200, height=200)
            image = DecodedStreamObject()
            image.set_data(bytes([255, 255, 255]))
            image.update({NameObject('/Type'): NameObject('/XObject'),
                          NameObject('/Subtype'): NameObject('/Image'),
                          NameObject('/Width'): NumberObject(1),
                          NameObject('/Height'): NumberObject(1),
                          NameObject('/ColorSpace'): NameObject('/DeviceRGB'),
                          NameObject('/BitsPerComponent'): NumberObject(8)})
            page[NameObject('/Resources')] = DictionaryObject({NameObject('/XObject'):
                DictionaryObject({NameObject('/Im1'): writer._add_object(image)})})
            content = DecodedStreamObject()
            content.set_data(b'q 10 0 0 10 0 0 cm /Im1 Do Q')
            page[NameObject('/Contents')] = writer._add_object(content)
            with (root / 'scan.pdf').open('wb') as out:
                writer.write(out)
            report = check_source_folder(root)
            self.assertFalse(report.passed)  # blank text is an ingestion error
            self.assertEqual(report.image_only_units, 1)
            self.assertIn('OCR may be needed', report.table())
            self.assertIn('likely_image_only_page', {g['code'] for g in report.by_code()})

    def test_cli_folder_only_and_detail_default(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / 'one.txt').write_text('Hello from a source file.')
            with mock.patch('sys.stdout') as output:
                self.assertEqual(main(['check', str(root), '--json']), 0)
            emitted = ''.join(call.args[0] for call in output.write.call_args_list if call.args)
            payload = json.loads(emitted)
            self.assertEqual(payload['files_audited'], 1)
            self.assertNotIn('detailed_audit', payload)
